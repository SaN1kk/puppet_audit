import csv
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import report_csv as report


class ReportTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.source = Path(self.temporary.name) / "hosts.csv"
        self.output = self.source.with_suffix(".xlsx")

    def write_csv(self, rows):
        fields = list(dict.fromkeys(key for row in rows for key in row))
        with self.source.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.DictWriter(file, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)

    def test_json_paths_types_and_empty_containers(self):
        self.write_csv([{
            "certname": "host-a", "uuid": "000001",
            "interfaces": json.dumps({"eth0.100": {"up": False, "count": 0}}),
            "cpu_models": '["Xeon", "EPYC"]', "disks": "{}", "partitions": "[]",
            "telegraf": '{"ensure": null}',
        }])
        rows, fields = report.read_csv(self.source)
        self.assertIn('interfaces["eth0.100"].up', fields)
        self.assertIs(rows[0]['interfaces["eth0.100"].up'], False)
        self.assertEqual(rows[0]['interfaces["eth0.100"].count'], 0)
        self.assertEqual(rows[0]["cpu_models[1]"], "EPYC")
        self.assertEqual(rows[0]["uuid"], "000001")
        self.assertEqual(rows[0]["disks"], "{}")
        self.assertEqual(rows[0]["partitions"], "[]")
        self.assertIsNone(rows[0]["telegraf.ensure"])

    def test_comparison_is_key_based_and_failed_hosts_are_excluded(self):
        self.write_csv([
            {"certname": "a", "uuid": "uuid-a", "telegraf": '{"enable": false, "port": 0}', "os": "Linux", "error": ""},
            {"certname": "b", "uuid": "uuid-b", "telegraf": '{"port": 0, "enable": false}', "os": "Linux", "error": ""},
            {"certname": "failed", "error": "HTTP 500"},
        ])
        rows, fields = report.read_csv(self.source)
        valid, selected, changed = report.differences(rows, fields, report.DEFAULT_COMPARE)
        self.assertEqual(len(valid), 2)
        self.assertIn("telegraf.enable", selected)
        self.assertNotIn("uuid", selected)
        self.assertEqual(changed, set())
        self.assertEqual(report.differences(rows, fields, ["uuid"])[2], {"uuid"})

    def test_false_and_zero_are_distinct(self):
        rows = [{"certname": "a", "value": False}, {"certname": "b", "value": 0}]
        self.assertEqual(report.differences(rows, ["value"], ["value"])[2], {"value"})
        self.assertFalse(report.missing(False))
        self.assertFalse(report.missing(0))

    def test_workbook_styles_missing_values_and_text(self):
        self.write_csv([
            {"certname": "host-a", "os_version": "22.04", "bios_version": "1.0", "uuid": "000001", "model": "=literal", "auditd": '{"enable": false, "counter": 12345678901234567890}', "error": ""},
            {"certname": "host-b", "os_version": "24.04", "bios_version": "", "uuid": "000002", "model": "=literal", "auditd": '{"enable": false, "counter": 12345678901234567890}', "error": ""},
            {"certname": "failed", "error": "HTTP 500"},
        ])
        original = self.source.read_bytes()
        rows, fields = report.read_csv(self.source)
        count = report.write_report(rows, fields, report.DEFAULT_COMPARE, self.output)
        self.assertEqual(count, 2)
        self.assertEqual(self.source.read_bytes(), original)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["Различия", "Хосты"])
        sheet = book["Хосты"]
        columns = {cell.value: cell.column for cell in sheet[5]}
        self.assertEqual(sheet.cell(6, columns["os_version"]).fill.fgColor.rgb[-6:], report.YELLOW)
        self.assertEqual(sheet.cell(7, columns["bios_version"]).fill.fgColor.rgb[-6:], report.RED)
        self.assertEqual(sheet.cell(8, columns["certname"]).fill.fgColor.rgb[-6:], report.GREY)
        self.assertEqual(sheet.cell(8, columns["error"]).fill.fgColor.rgb[-6:], report.RED)
        self.assertEqual(sheet.cell(6, columns["uuid"]).value, "000001")
        self.assertEqual(sheet.cell(6, columns["bios_version"]).value, "1.0")
        self.assertEqual(sheet.cell(6, columns["auditd.enable"]).value, "false")
        self.assertEqual(sheet.cell(6, columns["auditd.counter"]).value, "12345678901234567890")
        self.assertEqual(sheet.cell(6, columns["model"]).data_type, "s")
        self.assertEqual(sheet.freeze_panes, "B6")
        self.assertEqual(book["Различия"].auto_filter.ref, "A5:C7")

    def test_all_failed_hosts_do_not_look_like_a_successful_comparison(self):
        rows = [{"certname": "failed", "error": "HTTP 500"}]
        self.assertEqual(report.write_report(rows, list(rows[0]), ["*"], self.output), 0)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        self.assertIn("минимум два", book["Различия"]["A6"].value)
        self.assertEqual(book["Различия"].auto_filter.ref, "A5:A6")

    def test_invalid_input_and_cli_options(self):
        self.write_csv([{"certname": "same"}, {"certname": "same"}])
        with self.assertRaisesRegex(ValueError, "duplicate certname"):
            report.read_csv(self.source)
        self.write_csv([{"certname": "a", "disks": "{broken"}])
        with self.assertRaisesRegex(ValueError, "invalid JSON"):
            report.read_csv(self.source)
        self.write_csv([{"certname": "a", "os": "Linux"}])
        script = Path(report.__file__).resolve()
        command = [sys.executable, str(script), str(self.source)]
        result = subprocess.run(command + ["--compare", "typo"], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("No field matches", result.stderr)
        self.assertFalse(self.output.exists())
        result = subprocess.run(command + ["-o", str(self.source)], capture_output=True, text=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("separate .xlsx", result.stderr)

    def test_save_failure_keeps_existing_report(self):
        self.output.write_bytes(b"existing report")
        with patch("openpyxl.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                report.write_report([{"certname": "a"}], ["certname"], [], self.output)
        self.assertEqual(self.output.read_bytes(), b"existing report")
        self.assertEqual(list(self.output.parent.glob("*.xlsx")), [self.output])


if __name__ == "__main__":
    unittest.main()
