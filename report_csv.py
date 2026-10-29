#!/usr/bin/env python3
"""Expand JSON cells in a PuppetDB CSV and highlight differences in an Excel report."""

import argparse
import csv
import json
import math
import tempfile
from fnmatch import fnmatchcase
from pathlib import Path


DEFAULT_COMPARE = (
    "environment", "os", "os_version", "architecture", "kernel", "kernel_release",
    "manufacturer", "model", "bios_version", "bios_date", "board",
    "cpu_physical", "cpu_logical", "cpu_models*", "ram_total",
    "kesl*", "telegraf*", "zabbix_agent*", "auditd*",
)
YELLOW, RED, GREY, NAVY = "FFF2CC", "FCE4D6", "E7E6E6", "203864"


def flatten(value, path, result):
    if isinstance(value, dict) and value:
        for key, item in sorted(value.items()):
            suffix = f".{key}" if key.isidentifier() else f"[{json.dumps(key, ensure_ascii=False)}]"
            flatten(item, path + suffix, result)
    elif isinstance(value, list) and value:
        for index, item in enumerate(value):
            flatten(item, f"{path}[{index}]", result)
    else:
        result[path] = json.dumps(value) if isinstance(value, (dict, list)) else value


def read_csv(filename):
    csv.field_size_limit(10 * 1024 * 1024)
    rows, hosts = [], set()
    with open(filename, newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        if not reader.fieldnames or "certname" not in reader.fieldnames:
            raise ValueError("CSV must contain a certname column")
        if len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("CSV contains duplicate column names")
        for number, source in enumerate(reader, 2):
            if None in source or any(value is None for value in source.values()):
                raise ValueError(f"CSV row {number}: wrong number of columns")
            host = source["certname"].strip()
            if not host or host in hosts:
                raise ValueError(f"CSV row {number}: empty or duplicate certname")
            hosts.add(host)
            row = {}
            for column, value in source.items():
                if value.lstrip().startswith(("{", "[")):
                    try:
                        value = json.loads(value)
                    except json.JSONDecodeError as error:
                        raise ValueError(f"CSV row {number}, {column}: invalid JSON") from error
                flatten(value, column, row)
            row["certname"] = host
            rows.append(row)
    if not rows:
        raise ValueError("CSV contains no hosts")
    fields = list(dict.fromkeys(field for row in rows for field in row))
    return rows, fields


def missing(value):
    return value is None or value == ""


def differences(rows, fields, patterns):
    selected = {
        field for field in fields
        if any(field == pattern or fnmatchcase(field, pattern) for pattern in patterns)
    }
    valid = [row for row in rows if not row.get("error")]
    changed = set()
    for field in selected:
        values = {
            None if missing(row.get(field)) else json.dumps(row[field], ensure_ascii=False)
            for row in valid
        }
        if len(values) > 1:
            changed.add(field)
    return valid, selected, changed


def write_report(rows, fields, patterns, output):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    valid, selected, changed = differences(rows, fields, patterns)
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Различия"
    data = workbook.create_sheet("Хосты")
    widths = {}
    fills = {color: PatternFill("solid", fgColor=color) for color in (YELLOW, RED, GREY, NAVY)}

    def put(sheet, row, column, value, color=None):
        if isinstance(value, bool):
            value = str(value).lower()
        elif isinstance(value, int) and len(str(abs(value))) > 15:
            value = str(value)  # Excel numbers preserve only 15 significant digits.
        cell = sheet.cell(row, column, value)
        if isinstance(value, str):
            if len(value) > 32767:
                raise ValueError(f"Value exceeds Excel cell limit: {sheet.title}!{cell.coordinate}")
            cell.data_type = "s"  # CSV text, including '=' prefixes, stays text.
            cell.number_format = "@"
        cell.font = Font(name="Calibri", size=11, color="222222")
        cell.alignment = Alignment(horizontal="left" if isinstance(value, str) else "right", vertical="top", wrap_text=True)
        if color:
            cell.fill = fills[color]
        return cell

    def prepare(sheet, title, headers):
        if len(headers) > 16384:
            raise ValueError("Report exceeds Excel column limit; use a smaller host group")
        last = get_column_letter(len(headers))
        widths[sheet.title] = len(headers)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(2, min(8, len(headers))))
        put(sheet, 1, 1, title).font = Font(name="Calibri", size=16, bold=True, color=NAVY)
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, min(8, len(headers))))
        put(sheet, 2, 1, f"Хостов: {len(rows)}; ошибок сбора: {len(rows) - len(valid)}; различающихся полей: {len(changed)}")
        for column, (label, color) in enumerate((
            ("Различающиеся значения", YELLOW), ("Нет данных / ошибка", RED), ("Ошибка сбора: не сравнивается", GREY),
        ), 1):
            put(sheet, 3, column, label, color)
        for column, label in enumerate(headers, 1):
            put(sheet, 5, column, label, NAVY).font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            sheet.column_dimensions[get_column_letter(column)].width = max(24, min(36, len(label) + 2))
        sheet.freeze_panes = "B6"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 28
        sheet.row_dimensions[3].height = 32
        header_lines = max(math.ceil(len(label) / max(8, sheet.column_dimensions[get_column_letter(column)].width - 3))
                           for column, label in enumerate(headers, 1))
        sheet.row_dimensions[5].height = max(32, header_lines * 15 + 6)
        sheet.auto_filter.ref = f"A5:{last}5"

    def finish(sheet):
        if sheet.max_row > 1048576:
            raise ValueError("Report exceeds Excel row limit")
        if sheet.max_row > 5:
            sheet.auto_filter.ref = f"A5:{get_column_letter(widths[sheet.title])}{sheet.max_row}"
        for row in sheet.iter_rows(min_row=6, max_col=widths[sheet.title]):
            lines = max((
                sum(max(1, math.ceil(len(line) / max(8, sheet.column_dimensions[cell.column_letter].width - 3)))
                    for line in str(cell.value if cell.value is not None else "").split("\n"))
                for cell in row
            ), default=1)
            sheet.row_dimensions[row[0].row].height = min(409, max(24, lines * 15 + 6))

    prepare(summary, "Различия между хостами", ["Поле", *[row["certname"] for row in valid]])
    summary.column_dimensions["A"].width = 56
    for number, field in enumerate((field for field in fields if field in changed), 6):
        put(summary, number, 1, field)
        for column, row in enumerate(valid, 2):
            value = row.get(field)
            put(summary, number, column, value, RED if missing(value) else YELLOW)
    if not changed:
        message = "Различий по выбранным полям нет" if len(valid) >= 2 else "Для сравнения нужны минимум два хоста без ошибок"
        put(summary, 6, 1, message)

    prepare(data, "Данные хостов: JSON развёрнут по полям", fields)
    for number, row in enumerate(rows, 6):
        failed = bool(row.get("error"))
        for column, field in enumerate(fields, 1):
            value, color = row.get(field), None
            if failed:
                color = RED if field == "error" else GREY
            elif field in selected and missing(value):
                color = RED
            elif field in changed:
                color = YELLOW
            put(data, number, column, value, color)
    for sheet in workbook:
        finish(sheet)
    # Finish the new workbook before replacing an existing report.
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix=".xlsx", delete=False) as file:
        temporary = Path(file.name)
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return len(changed)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("csv_file", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    parser.add_argument("--compare", nargs="+", metavar="FIELD", help="Columns or glob patterns; quote '*' in the shell")
    args = parser.parse_args()
    output = args.output or args.csv_file.with_suffix(".xlsx")
    if output.resolve() == args.csv_file.resolve() or output.suffix.lower() != ".xlsx":
        parser.error("Output must be a separate .xlsx file")
    try:
        rows, fields = read_csv(args.csv_file)
        if args.compare:
            for pattern in args.compare:
                if not any(field == pattern or fnmatchcase(field, pattern) for field in fields):
                    raise ValueError(f"No field matches --compare {pattern!r}")
        count = write_report(rows, fields, args.compare or DEFAULT_COMPARE, output)
    except ImportError:
        parser.exit(1, "Install report dependencies: python3 -m pip install -r requirements-report.txt\n")
    except (OSError, ValueError, csv.Error) as error:
        parser.exit(1, f"Error: {error}\n")
    print(f"Saved: {output} | Hosts: {len(rows)} | Differing fields: {count}")


if __name__ == "__main__":
    main()
